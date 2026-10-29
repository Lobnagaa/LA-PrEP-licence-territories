"""Turn the public spreadsheet into the JSON the page reads.

Usage (from the repository folder):
    python3 tools/xlsx_to_json.py [path/to/lapal_licence_coverage.xlsx]

Reads   ../licence-territory-dashboard/lapal-public-handoff/data/lapal_licence_coverage.xlsx (default), read-only
Writes  data/lapal_licence_coverage.json

This is the only step between the spreadsheet and the page. It reshapes, it does
not compute: no coverage figure is precomputed here (the page computes them live
from the Coverage sheet). Content checks (unknown ISO3, missing licence row,
unknown basis...) live in src/data.js so the page and this build apply the same
rules; run `node tools/check_data.js` after this script.

Spreadsheet contract (sheet -> one JSON key):
    Countries  -> countries[]   one row per country, keyed by iso3
    Licences   -> licences[]    one row per licence, keyed by code; row order = display order
    Coverage   -> coverage[]    one row per country x licence
    Notes      -> generated, impact{}, texts{}   key/value rows:
                  generated, impact_year, impact_global_new_infections,
                  impact_source, text_<name> (each becomes texts.<name>)
Columns are read by header name, so their order in the sheet does not matter.
Extra columns are carried through; missing required columns stop the build.
"""
import json
import os
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = os.path.join(ROOT, '..', 'licence-territory-dashboard', 'lapal-public-handoff', 'data', 'lapal_licence_coverage.xlsx')
OUT = os.path.join(ROOT, 'data', 'lapal_licence_coverage.json')

REQUIRED = {
    'Countries': ['iso3', 'country', 'region', 'income_group', 'income_bucket',
                  'new_hiv_infections', 'new_hiv_infections_shown_as',
                  'hiv_incidence_per_1000_uninfected'],
    'Licences': ['code', 'name', 'licence_type', 'originator', 'disease', 'indication',
                 'colour', 'fade_colour', 'ramp_direct', 'ramp_indirect'],
    'Coverage': ['iso3', 'licence', 'covered', 'basis', 'royalty_rate'],
    'Notes': ['key', 'value'],
}
# Numeric columns: the page expects numbers (or null), never text.
NUMERIC = {'new_hiv_infections', 'hiv_incidence_per_1000_uninfected'}


def clean(value):
    """Blank cells and whitespace-only strings become null; strings are trimmed."""
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def read_sheet(wb, name):
    if name not in wb.sheetnames:
        sys.exit(f'ERROR: sheet {name!r} is missing from the spreadsheet')
    rows = list(wb[name].iter_rows(values_only=True))
    header = [clean(h) for h in rows[0]]
    missing = [c for c in REQUIRED[name] if c not in header]
    if missing:
        sys.exit(f'ERROR: sheet {name!r} is missing column(s) {missing}')
    out = []
    for n, row in enumerate(rows[1:], start=2):
        if all(clean(v) is None for v in row):
            continue  # skip fully blank rows (common at the end of edited sheets)
        rec = {h: clean(v) for h, v in zip(header, row) if h}
        for col in NUMERIC & rec.keys():
            v = rec[col]
            if v is not None and not isinstance(v, (int, float)):
                sys.exit(f'ERROR: {name} row {n}: {col} must be a number or blank, got {v!r}')
            if v is not None:
                rec[col] = float(v)
        out.append(rec)
    return out


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = load_workbook(src, read_only=True, data_only=True)

    notes = {r['key']: r['value'] for r in read_sheet(wb, 'Notes')}
    for key in ('generated', 'impact_year', 'impact_global_new_infections', 'impact_source'):
        if notes.get(key) is None:
            sys.exit(f'ERROR: Notes sheet has no value for {key!r}')

    payload = {
        'generated': str(notes['generated'])[:10],
        'licences': read_sheet(wb, 'Licences'),
        'countries': read_sheet(wb, 'Countries'),
        'coverage': read_sheet(wb, 'Coverage'),
        'texts': {k[len('text_'):]: v for k, v in notes.items() if k.startswith('text_')},
        'impact': {
            'year': str(notes['impact_year']),
            'global_new_infections': float(notes['impact_global_new_infections']),
            'source': notes['impact_source'],
        },
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)
    print(f'wrote {os.path.relpath(OUT, ROOT)}: {len(payload["countries"])} countries, '
          f'{len(payload["licences"])} licences, {len(payload["coverage"])} coverage rows')


if __name__ == '__main__':
    main()
